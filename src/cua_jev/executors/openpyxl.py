from __future__ import annotations

from typing import Any

from ..errors import CapabilityUnavailable
from ..models import ActionCandidate, ActionReceipt
from .common import execute_with_receipt


class OpenPyxlExecutor:
    """Direct workbook-file operations that do not require a running Excel process."""

    def __call__(self, candidate: ActionCandidate, observation_id: str, decision_id: str) -> ActionReceipt:
        def operation() -> dict[str, Any]:
            try:
                from openpyxl import load_workbook
                from openpyxl.chart import BarChart, Reference
            except ImportError:
                raise CapabilityUnavailable("install cua-jev[windows] for workbook file APIs") from None

            args = candidate.arguments
            workbook = load_workbook(args["workbook_path"])
            try:
                sheet = workbook[args["sheet"]]
                if candidate.capability == "excel.file_write_formula":
                    sheet[args["cell"]] = args["formula"]
                elif candidate.capability == "excel.file_write_value":
                    sheet[args["cell"]] = args["value"]
                elif candidate.capability == "excel.file_create_chart":
                    chart = BarChart()
                    chart.title = args.get("title", "Chart")
                    chart.y_axis.title = args.get("y_axis_title", "Value")
                    chart.x_axis.title = args.get("x_axis_title", "Category")
                    values = Reference(
                        sheet,
                        min_col=int(args.get("value_column", 2)),
                        min_row=1,
                        max_row=int(args["max_row"]),
                    )
                    categories = Reference(
                        sheet,
                        min_col=int(args.get("category_column", 1)),
                        min_row=2,
                        max_row=int(args["max_row"]),
                    )
                    chart.add_data(values, titles_from_data=True)
                    chart.set_categories(categories)
                    sheet.add_chart(chart, args.get("anchor", "D2"))
                else:
                    raise ValueError(f"unsupported workbook-file capability: {candidate.capability}")
                workbook.save(args["workbook_path"])
            finally:
                workbook.close()
            return {
                "workbook_path": args["workbook_path"],
                "sheet": args["sheet"],
                "operation": candidate.capability,
            }

        return execute_with_receipt(candidate, observation_id, decision_id, operation)
